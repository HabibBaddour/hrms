import csv
from datetime import datetime
from io import BytesIO, StringIO

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from departments.models import Department
from .models import LeaveRequest
from employees.models import Employee
from .ml_engine import is_peak_period, predict_leave_approval
from .ml.predictor import predict_leave
from .services import deduplicate_leave_queryset, get_department_manager, notify_leave_status_changed, notify_leave_submitted

def _employee_search_query(search_query):
    q = Q()
    if not search_query:
        return q

    search_value = search_query.strip()
    if not search_value:
        return q

    if search_value.isdigit():
        q |= Q(employee__id=search_value)
        q |= Q(id=int(search_value))
    else:
        q |= Q(employee__first_name__icontains=search_value)
        q |= Q(employee__last_name__icontains=search_value)
        q |= Q(employee__user__first_name__icontains=search_value)
        q |= Q(employee__user__last_name__icontains=search_value)
        q |= Q(employee__user__username__icontains=search_value)
        q |= Q(employee__national_id__icontains=search_value)

    return q


def _render_leave_requests_excel(leaves):
    """Build a real .xlsx spreadsheet from a filtered leave request queryset."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'طلبات الإجازات'

    headers = [
        'رقم الطلب', 'الموظف', 'رقم الموظف', 'القسم', 'المسمى الوظيفي',
        'نوع الإجازة', 'تاريخ البداية', 'تاريخ النهاية', 'عدد الأيام',
        'الحالة', 'توصية الذكاء الاصطناعي', 'نسبة الثقة', 'تاريخ التقديم',
    ]
    header_fill = PatternFill(start_color='4A3AB8', end_color='4A3AB8', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    status_labels = {
        'PENDING': 'معلق',
        'APPROVED': 'مقبول',
        'REJECTED': 'مرفوض',
    }
    ai_labels = {
        'APPROVED': 'موصى بالقبول',
        'REJECTED': 'موصى بالرفض',
        'PENDING': 'قيد التحليل',
    }

    for row_index, leave in enumerate(leaves, start=2):
        department = getattr(leave.employee.department, 'name', '-') if leave.employee.department else '-'
        position = getattr(leave.employee.position, 'title', '') if leave.employee.position else ''
        row = [
            leave.pk,
            leave.employee.get_full_name(),
            getattr(leave.employee, 'employee_number', '') or '',
            department,
            position or '-',
            leave.get_leave_type_display(),
            leave.start_date.strftime('%Y-%m-%d'),
            leave.end_date.strftime('%Y-%m-%d'),
            leave.total_days,
            status_labels.get(leave.status, leave.status),
            ai_labels.get(leave.ai_prediction, leave.ai_prediction),
            f'{leave.ai_confidence}%',
            leave.created_at.strftime('%Y-%m-%d %H:%M') if leave.created_at else '',
        ]
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column, value=value)

    for column in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 16
    sheet.freeze_panes = 'A2'

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="leave_requests.xlsx"'
    return response


def _leave_list_context(request):
    """تجمع سياق قائمة طلبات الإجازات (يُستخدم في صفحة الإجازات نفسها
    وفي تبويب 'الإجازات' المدمج داخل صفحة جدول الدوام والإجازات)."""
    employee_profile = getattr(request.user, 'employee_profile', None)
    position_role = getattr(getattr(employee_profile, 'position', None), 'role', '')
    can_view_all = (
        request.user.is_superuser
        or request.user.is_staff
        or position_role.lower() == 'hr admin'
        or request.user.groups.filter(name='HR').exists()
    )
    is_department_manager = bool(
        employee_profile and employee_profile.department_id and
        getattr(employee_profile.position, 'role', '') == 'Manager'
    )

    dept_id = request.GET.get('department')
    role = request.GET.get('role')
    leave_type = request.GET.get('leave_type')
    status = request.GET.get('status')
    search_query = request.GET.get('q', '').strip()

    departments = Department.objects.order_by('name')
    role_choices = [
        ('Employee', 'موظف'),
        ('Manager', 'مدير'),
        ('HR Admin', 'مسؤول موارد بشرية'),
    ]

    if can_view_all or is_department_manager:
        leaves = LeaveRequest.objects.select_related('employee__user', 'employee__position__department').all()
        employees = Employee.objects.select_related('user', 'department', 'position').all()

        if is_department_manager and not can_view_all:
            leaves = leaves.filter(employee__department_id=employee_profile.department_id)
            employees = employees.filter(department_id=employee_profile.department_id)

        if dept_id:
            leaves = leaves.filter(employee__department_id=dept_id)
            employees = employees.filter(department_id=dept_id)
        if role:
            leaves = leaves.filter(employee__position__role=role)
            employees = employees.filter(position__role=role)
        if leave_type:
            leaves = leaves.filter(leave_type=leave_type)
        if search_query:
            search_filter = Q()
            if search_query.isdigit():
                search_filter |= Q(employee__id=int(search_query))
                search_filter |= Q(id=int(search_query))
            else:
                search_filter |= Q(employee__first_name__icontains=search_query)
                search_filter |= Q(employee__last_name__icontains=search_query)
                search_filter |= Q(employee__user__first_name__icontains=search_query)
                search_filter |= Q(employee__user__last_name__icontains=search_query)
                search_filter |= Q(employee__user__username__icontains=search_query)
                search_filter |= Q(employee__national_id__icontains=search_query)
                search_filter |= Q(first_name__icontains=search_query)
                search_filter |= Q(last_name__icontains=search_query)
                search_filter |= Q(user__first_name__icontains=search_query)
                search_filter |= Q(user__last_name__icontains=search_query)
                search_filter |= Q(user__username__icontains=search_query)
                search_filter |= Q(national_id__icontains=search_query)
            leaves = leaves.filter(search_filter)
            employees = employees.filter(search_filter)
        if status:
            if status == 'RECOMMENDED_AI':
                leaves = leaves.filter(ai_prediction__in=['APPROVED', 'REJECTED'])
            else:
                leaves = leaves.filter(status=status)

        employees = employees.prefetch_related('leave_requests')
    else:
        leaves = LeaveRequest.objects.filter(employee__user=request.user)
        employees = []

        if dept_id:
            leaves = leaves.filter(employee__department_id=dept_id)
        if role:
            leaves = leaves.filter(employee__position__role=role)
        if leave_type:
            leaves = leaves.filter(leave_type=leave_type)
        if search_query:
            search_filter = Q()
            if search_query.isdigit():
                search_filter |= Q(employee__id=int(search_query))
            else:
                search_filter |= Q(employee__first_name__icontains=search_query)
                search_filter |= Q(employee__last_name__icontains=search_query)
                search_filter |= Q(employee__user__first_name__icontains=search_query)
                search_filter |= Q(employee__user__last_name__icontains=search_query)
                search_filter |= Q(employee__user__username__icontains=search_query)
                search_filter |= Q(employee__national_id__icontains=search_query)
            leaves = leaves.filter(search_filter)
        if status:
            if status == 'RECOMMENDED_AI':
                leaves = leaves.filter(ai_prediction__in=['APPROVED', 'REJECTED'])
            else:
                leaves = leaves.filter(status=status)

    leaves = deduplicate_leave_queryset(leaves)

    return {
        'leaves': leaves,
        'filtered_count': leaves.count(),
        'pending_count': leaves.filter(status='PENDING').count(),
        'approved_count': leaves.filter(status='APPROVED').count(),
        'rejected_count': leaves.filter(status='REJECTED').count(),
        'employees': employees,
        'can_view_all': can_view_all,
        'is_privileged': can_view_all or is_department_manager,
        'departments': departments,
        'roles': role_choices,
        'selected_department': dept_id,
        'selected_role': role,
        'selected_leave_type': leave_type,
        'selected_status': status,
        'search_query': search_query,
    }


@login_required
def leave_list(request):
    context = _leave_list_context(request)
    leaves = context['leaves']
    employees = context['employees']
    can_view_all = context['can_view_all']

    export_balances = request.GET.get('export') == 'excel'
    export_leaves = request.GET.get('export') == 'leaves'

    if export_leaves:
        return _render_leave_requests_excel(leaves)

    if export_balances and can_view_all:
        csv_buffer = StringIO()
        writer = csv.writer(csv_buffer)
        writer.writerow(['Employee ID', 'Full Name', 'Department', 'Role', 'Annual Balance', 'Total Annual Days'])

        for employee in employees:
            department_name = employee.department.name if employee.department else '-'
            role_name = employee.position.role if employee.position else '-'
            balance = employee.get_annual_leave_balance()
            writer.writerow([
                employee.employee_number,
                employee.get_full_name(),
                department_name,
                role_name,
                balance,
                employee.ANNUAL_LEAVE_DAYS,
            ])

        response = HttpResponse(csv_buffer.getvalue(), content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="employee_leave_balances.csv"'
        return response

    context['export_leaves'] = export_leaves
    return render(request, 'leaves/leave_list.html', context)


def _can_review_leave(request, leave):
    """هل لدى المستخدم صلاحية رؤية/مراجعة هذا الطلب؟"""
    employee_profile = getattr(request.user, 'employee_profile', None)
    is_hr = (
        request.user.is_superuser or request.user.is_staff or
        getattr(getattr(employee_profile, 'position', None), 'role', '').lower() == 'hr admin' or
        request.user.groups.filter(name='HR').exists()
    )
    if is_hr:
        return True
    is_department_manager = bool(
        employee_profile and employee_profile.department_id == leave.employee.department_id and
        getattr(employee_profile.position, 'role', '') == 'Manager'
    )
    if is_department_manager:
        return True
    # الموظف يرى طلبه الخاص فقط
    return leave.employee.user_id == request.user.id


@login_required
def leave_detail_view(request, leave_id):
    """عرض تفاصيل طلب إجازة. يدعم استدعاء AJAX (JSON) للمودال أو صفحة كاملة."""
    leave = get_object_or_404(
        LeaveRequest.objects.select_related('employee__user', 'employee__position__department'),
        pk=leave_id,
    )
    if not _can_review_leave(request, leave):
        messages.error(request, 'ليس لديك صلاحية الاطلاع على هذا الطلب.')
        return redirect('leaves:leave_list')

    data = {
        'id': leave.pk,
        'employee_name': leave.employee.get_full_name(),
        'employee_id': leave.employee_id,
        'leave_type': leave.get_leave_type_display(),
        'leave_type_code': leave.leave_type,
        'start_date': str(leave.start_date),
        'end_date': str(leave.end_date),
        'total_days': leave.total_days,
        'reason': leave.reason,
        'status': leave.get_status_display(),
        'status_code': leave.status,
        'ai_prediction': leave.get_ai_prediction_display() if hasattr(leave, 'get_ai_prediction_display') else leave.ai_prediction,
        'ai_prediction_code': leave.ai_prediction,
        'ai_confidence': leave.ai_confidence,
        'attachment_url': leave.attachment.url if leave.attachment else None,
        'created_at': leave.created_at.strftime('%Y-%m-%d %H:%M') if leave.created_at else None,
        'manager_notes': leave.manager_notes,
        'approved_by': leave.approved_by.get_full_name() if leave.approved_by else None,
    }

    if request.headers.get('x-requested-with') == 'XMLHttpRequest' or request.GET.get('format') == 'json':
        return JsonResponse({'success': True, 'leave': data})

    return render(request, 'leaves/leave_detail.html', {
        'leave': leave,
        'leave_data': data,
    })


@login_required
def leave_delete_view(request, leave_id):
    """حذف طلب إجازة. الموظف يحذف طلبه المعلّق فقط، ومسؤولو HR يحذفون أي طلب."""
    leave = get_object_or_404(LeaveRequest, pk=leave_id)
    employee_profile = getattr(request.user, 'employee_profile', None)
    is_hr = (
        request.user.is_superuser or request.user.is_staff or
        getattr(getattr(employee_profile, 'position', None), 'role', '').lower() == 'hr admin' or
        request.user.groups.filter(name='HR').exists()
    )

    is_owner = leave.employee.user_id == request.user.id
    if not (is_hr or is_owner):
        messages.error(request, 'ليس لديك صلاحية حذف هذا الطلب.')
        return redirect('leaves:leave_list')

    # الموظف يحذف طلبه المعلّق فقط؛ مسؤولو HR يحذفون أي حالة
    if is_owner and not is_hr and leave.status != 'PENDING':
        messages.error(request, 'لا يمكن حذف طلب إجازة مكتمل المعالجة (مقبول/مرفوض).')
        return redirect('leaves:leave_list')

    if request.method == 'POST':
        leave.delete()
        messages.success(request, 'تم حذف طلب الإجازة بنجاح.')
        return redirect('leaves:leave_list')

    return render(request, 'leaves/leave_confirm_delete.html', {'leave': leave})


@login_required
def analyze_leave_ai(request):
    """تحليل ذكي مبدئي لطلب إجازة (AJAX/POST) — لا يُنشئ أي طلب."""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'الطريقة غير مسموحة.'}, status=405)

    employee = getattr(request.user, 'employee_profile', None)
    if not employee:
        return JsonResponse({'success': False, 'error': 'يجب ربط حسابك بملف موظف أولاً.'}, status=400)

    leave_type = request.POST.get('leave_type', 'ANNUAL')
    if leave_type not in dict(LeaveRequest.LeaveType.choices):
        return JsonResponse({'success': False, 'error': 'نوع الإجازة غير صالح.'}, status=400)

    try:
        start_date = datetime.strptime(request.POST.get('start_date', ''), '%Y-%m-%d').date()
        end_date = datetime.strptime(request.POST.get('end_date', ''), '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return JsonResponse({'success': False, 'error': 'يرجى إدخال تاريخين صحيحين للبداية والنهاية.'}, status=400)

    if end_date < start_date:
        return JsonResponse({'success': False, 'error': 'يجب أن يكون تاريخ النهاية بعد تاريخ البداية.'}, status=400)

    requested_days = (end_date - start_date).days + 1
    remaining_balance = employee.leave_remaining(leave_type)

    dept_conflicts = 0
    if employee.department_id:
        dept_conflicts = LeaveRequest.objects.filter(
            employee__department_id=employee.department_id,
            status='APPROVED',
            start_date__lte=end_date,
            end_date__gte=start_date,
        ).exclude(employee=employee).count()

    is_peak = is_peak_period(start_date, end_date)
    notice_days = (start_date - timezone.localdate()).days

    result = predict_leave_approval(
        requested_days=requested_days,
        remaining_balance=remaining_balance,
        dept_conflicts=dept_conflicts,
        is_peak=is_peak,
        notice_days=notice_days,
        leave_type=leave_type,
    )

    return JsonResponse({
        'success': True,
        'leave_type': leave_type,
        'requested_days': requested_days,
        'remaining_balance': remaining_balance,
        'dept_conflicts': dept_conflicts,
        'is_peak': is_peak,
        'notice_days': max(notice_days, 0),
        'probability': result['probability'],
        'verdict': result['verdict'],
        'recommendation': result['recommendation'],
        'reasons': result['reasons'],
        'model': result['model'],
    })


def _apply_leave_context(request):
    """سياق نموذج تقديم طلب إجازة (يُستخدم في صفحته المستقلة والتبويب المدمج)."""
    employee = getattr(request.user, 'employee_profile', None)
    recent_leaves = []
    if employee:
        recent_leaves = list(
            LeaveRequest.objects.filter(employee=employee)
            .order_by('-created_at')[:4]
        )
    return {
        'employee': employee,
        'today': timezone.localdate(),
        'recent_leaves': recent_leaves,
        'annual_left': employee.leave_remaining('ANNUAL') if employee else None,
        'sick_left': employee.leave_remaining('SICK') if employee else None,
        'emergency_left': employee.leave_remaining('EMERGENCY') if employee else None,
        'annual_quota': employee.leave_quota('ANNUAL') if employee else None,
        'sick_quota': employee.leave_quota('SICK') if employee else None,
        'emergency_quota': employee.leave_quota('EMERGENCY') if employee else None,
        'next': _safe_next_url(request),
    }


def _safe_next_url(request):
    """العودة الآمنة (نفس المضيف) إلى صفحة المصدر بعد إرسال الطلب."""
    next_url = request.POST.get('next') or request.GET.get('next')
    if next_url and url_has_allowed_host_and_scheme(
            next_url, allowed_hosts={request.get_host()}, require_https=request.is_secure()):
        return next_url
    return None


@login_required
def apply_leave(request):
    if request.method == 'POST':
        leave_type = request.POST.get('leave_type')
        start_date = request.POST.get('start_date')
        end_date = request.POST.get('end_date')
        reason = request.POST.get('reason')
        attachment = request.FILES.get('attachment')

        try:
            start_date = datetime.strptime(start_date, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_date, '%Y-%m-%d').date()
        except (TypeError, ValueError):
            messages.error(request, '???? ????? ?????? ????? ?????? ??????.')
            return render(request, 'leaves/apply_leave.html', _apply_leave_context(request))

        if end_date < start_date:
            messages.error(request, 'يجب أن يكون تاريخ النهاية بعد تاريخ البداية.')
            return render(request, 'leaves/apply_leave.html', _apply_leave_context(request))

        if attachment:
            allowed = ('.pdf', '.jpg', '.jpeg', '.png')
            ext = getattr(attachment, 'name', '') or ''
            if not ext.lower().endswith(allowed):
                messages.error(request, 'صيغة المرفق غير مدعومة. يُسمح فقط بملفات PDF أو JPG أو PNG.')
                return render(request, 'leaves/apply_leave.html', _apply_leave_context(request))

        employee = getattr(request.user, 'employee_profile', None)

        if employee:
            overlapping_request = LeaveRequest.objects.filter(
                employee=employee,
                status__in=['PENDING', 'APPROVED'],
                start_date__lte=end_date,
                end_date__gte=start_date,
            ).exists()
            if overlapping_request:
                messages.error(request, 'يوجد طلب إجازة مسجل بالفعل لهذا الموظف خلال هذه الفترة.')
                return render(request, 'leaves/apply_leave.html', _apply_leave_context(request))

            leave = LeaveRequest.objects.create(
                employee=employee,
                leave_type=leave_type,
                start_date=start_date,
                end_date=end_date,
                reason=reason,
                attachment=attachment if attachment else None,
                status='PENDING'
            )
            predict_leave_status(leave)
            notify_leave_submitted(leave, actor=request.user)
            messages.success(request, 'تم تقديم طلب الإجازة بنجاح. سيتم مراجعته من قبل الإدارة.')
            return redirect(_safe_next_url(request) or 'leaves:leave_list')

        messages.error(request, 'يجب ربط حسابك بملف موظف قبل تقديم طلب الإجازة.')

    return render(request, 'leaves/apply_leave.html', _apply_leave_context(request))


@login_required
def _can_decide_leave(request, leave):
    """هل يملك المستخدم صلاحية قبول/رفض هذا الطلب؟ فقط مدير قسم الموظف صاحب الطلب."""
    employee_profile = getattr(request.user, 'employee_profile', None)
    if not employee_profile:
        return False
    department = leave.employee.department or getattr(leave.employee.position, 'department', None)
    manager = get_department_manager(department)
    return bool(manager and manager.pk == employee_profile.pk)


def _decision_redirect(request):
    """العودة إلى صفحة الطلب (للمدير) بعد اتخاذ القرار، مع منع إعادة التوجيه لمواقع خارجية."""
    next_url = request.GET.get('next') or request.POST.get('next')
    if next_url and url_has_allowed_host_and_scheme(
            next_url, allowed_hosts={request.get_host()}, require_https=request.is_secure()):
        return redirect(next_url)
    return redirect('leaves:leave_list')


def _apply_leave_decision(request, leave, decision):
    """تطبيق قرار القبول/الرفض على طلب إجازة مع تفعيل الرصيد تلقائياً للمقبول."""
    if not _can_decide_leave(request, leave):
        messages.error(request, 'ليس لديك صلاحية مراجعة طلب الإجازة هذا.')
        return _decision_redirect(request)

    if decision not in ('APPROVED', 'REJECTED'):
        messages.error(request, 'اختيار القرار غير صالح.')
        return _decision_redirect(request)

    leave.status = decision
    leave.approved_by = getattr(request.user, 'employee_profile', None)
    leave.manager_notes = request.POST.get('manager_notes', '')
    leave.save(update_fields=['status', 'approved_by', 'manager_notes'])
    notify_leave_status_changed(leave, actor=request.user)
    if decision == 'APPROVED':
        messages.success(request, 'تم قبول طلب الإجازة وتحديث رصيد الموظف تلقائياً.')
    else:
        messages.success(request, 'تم رفض طلب الإجازة.')
    return _decision_redirect(request)


@login_required
def approve_leave(request, pk):
    leave = get_object_or_404(LeaveRequest, pk=pk)

    balance_total = leave.employee.leave_quota(leave.leave_type)
    balance_remaining = leave.employee.leave_remaining(leave.leave_type)
    balance_used = max(balance_total - balance_remaining, 0)

    if request.method == 'POST':
        return _apply_leave_decision(request, leave, request.POST.get('decision'))

    return render(request, 'leaves/approve_leave.html', {
        'leave': leave,
        'balance_total': balance_total,
        'balance_used': balance_used,
        'balance_remaining': balance_remaining,
    })


@login_required
def reject_leave(request, pk):
    """رفض طلب إجازة مباشرة من قائمة الطلبات (POST فقط)."""
    if request.method != 'POST':
        messages.error(request, 'الطريقة غير مسموحة.')
        return redirect('leaves:leave_list')
    leave = get_object_or_404(LeaveRequest, pk=pk)
    return _apply_leave_decision(request, leave, 'REJECTED')


def predict_leave_status(leave_instance):
    """
    توليد توصية الذكاء الاصطناعي باستخدام Random Forest
    لطلب إجازة مسجّل، مع حفظ التوصية ونسبة الثقة.
    """
    try:
        result = predict_leave(leave_instance)

        leave_instance.ai_prediction = result['prediction']
        leave_instance.ai_confidence = result['confidence']

        leave_instance.save(
            update_fields=['ai_prediction', 'ai_confidence']
        )

    except Exception as exc:
        print(f"ML Prediction Error: {exc}")